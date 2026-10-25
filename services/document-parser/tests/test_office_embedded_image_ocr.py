"""Exercise the real parser with text that exists only inside an Office image."""

from pathlib import Path

import pytest
from PIL import Image, ImageDraw, ImageFont

from app.services.document_parsing_service import parse_document
from app.services.rag_chunk_optimizer import build_rag_chunks


MARKER = "OFFICEPIC74219"


@pytest.fixture
def text_image(tmp_path: Path) -> Path:
    image = Image.new("RGB", (1800, 380), "white")
    font_path = next(path for path in (
        Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ) if path.exists())
    font = ImageFont.truetype(str(font_path), 130)
    ImageDraw.Draw(image).text((60, 90), MARKER, font=font, fill="black")
    path = tmp_path / "text.png"
    image.save(path)
    return path


def _make_office_file(tmp_path: Path, image: Path, ext: str) -> Path:
    path = tmp_path / f"embedded.{ext}"
    if ext == "docx":
        from docx import Document
        from docx.shared import Inches

        document = Document()
        document.add_paragraph("Ordinary document text")
        document.add_picture(str(image), width=Inches(6))
        document.save(path)
    elif ext == "pptx":
        from pptx import Presentation
        from pptx.util import Inches

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_picture(str(image), Inches(0.5), Inches(1), width=Inches(9))
        presentation.save(path)
    else:
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as WorkbookImage

        workbook = Workbook()
        workbook.active["A1"] = "Ordinary sheet text"
        workbook.active.add_image(WorkbookImage(str(image)), "B2")
        workbook.save(path)
    return path


@pytest.mark.parametrize("ext", ["docx", "pptx", "xlsx"])
def test_embedded_image_text_reaches_searchable_chunks(tmp_path: Path, text_image: Path, ext: str) -> None:
    source = _make_office_file(tmp_path, text_image, ext)
    before = source.read_bytes()

    parsed = parse_document(source, source.name)
    image_chunks = [
        chunk for chunk in parsed.raw_chunks or []
        if chunk.get("metadata", {}).get("sourceKind") == "office_embedded_image_ocr"
    ]
    rag = build_rag_chunks(parsed.raw_chunks or [])

    assert any(MARKER in chunk["text"] for chunk in image_chunks)
    assert any(MARKER in chunk["text"] for chunk in rag)
    assert parsed.raw["movoOfficeImageOcr"]["recognizedImageCount"] == 1
    assert source.read_bytes() == before
    if ext == "pptx":
        assert image_chunks[0]["pageNo"] == 1
    if ext == "xlsx":
        assert "xl/worksheets/sheet1.xml" in image_chunks[0]["metadata"]["locations"]


def test_existing_plain_text_stays_available(tmp_path: Path, text_image: Path) -> None:
    source = _make_office_file(tmp_path, text_image, "docx")
    parsed = parse_document(source, source.name)

    assert any("Ordinary document text" in chunk["text"] for chunk in parsed.raw_chunks or [])
