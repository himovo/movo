"""Read workbook data and provenance without changing the uploaded source."""

from __future__ import annotations

import tempfile
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile


_SHEET_TAG = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheet"


def normalized_workbook_path(source: Path, directory: Path) -> tuple[Path, list[str]]:
    """Normalize the known nonstandard OOXML sheet state in a temporary copy."""
    with ZipFile(source) as archive:
        xml = archive.read("xl/workbook.xml")
        root = ElementTree.fromstring(xml)
        repaired = []
        for sheet in root.iter(_SHEET_TAG):
            if sheet.get("state") == "show":
                sheet.set("state", "visible")
                repaired.append(str(sheet.get("name") or ""))
        if not repaired:
            return source, []
        target = directory / f"normalized-{source.name}"
        with ZipFile(target, "w") as output:
            for entry in archive.infolist():
                content = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True) if entry.filename == "xl/workbook.xml" else archive.read(entry)
                output.writestr(entry, content)
        return target, repaired


def parse_spreadsheet(path: Path):
    from openpyxl import load_workbook

    from app.services.document_parsing_service import ParsedDocument

    with tempfile.TemporaryDirectory(prefix="movo-workbook-") as temp:
        source, repaired = normalized_workbook_path(path, Path(temp))
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            chunks = []
            markdown_parts = []
            sheets = []
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)
                header_row = next(rows, ())
                headers = [str(value).strip() if value is not None else "" for value in header_row]
                visible_headers = [value for value in headers if value]
                sheets.append({"name": sheet.title, "headers": visible_headers, "state": sheet.sheet_state})
                heading = f"工作表：{sheet.title}\n字段：{'、'.join(visible_headers)}"
                markdown_parts.append(heading)
                chunks.append(_chunk(len(chunks), sheet.title, 1, heading, "sheet_summary"))
                for row_number, values in enumerate(rows, start=2):
                    cells = []
                    for index, value in enumerate(values):
                        if value is None or str(value).strip() == "":
                            continue
                        label = headers[index] if index < len(headers) and headers[index] else f"第{index + 1}列"
                        cells.append(f"{label}：{value}")
                    if not cells:
                        continue
                    text = f"工作表：{sheet.title}；第{row_number}行\n" + "；".join(cells)
                    markdown_parts.append(text)
                    chunks.append(_chunk(len(chunks), sheet.title, row_number, text, "table_row"))
            return ParsedDocument(
                markdown="\n\n".join(markdown_parts),
                raw={"parser": "openpyxl", "sheets": sheets, "normalizedSheetStates": repaired},
                raw_chunks=chunks,
            )
        finally:
            workbook.close()


def _chunk(index: int, sheet: str, row: int, text: str, kind: str) -> dict:
    return {
        "chunkId": f"raw_{index + 1:06d}", "ordinal": index,
        "text": text, "contextualText": text, "titlePath": [sheet],
        "pageNo": None, "contentType": kind, "chunkStage": "raw",
        "sourceChunkIds": [],
        "metadata": {"parser": "openpyxl", "sheetName": sheet, "rowNumber": row},
    }
